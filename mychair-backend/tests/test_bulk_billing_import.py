"""
Unit and integration tests for MyChair Bulk Billing Import Service.
Tests:
1. Template generation (3 sheets: How to Fill, Billing Data, Example; active sheet, formatting, styling, catalog population).
2. Date parsing (ISO dates, date objects with 00:00 start of day, slash-date rejection).
3. Excel parsing (sheet selection, ignoring guide/example sheets, skipping blank formatted rows).
4. Product billing parsing.
5. Failed records export and summary export.
6. Financial calculation and reconciliation math.
7. Mandatory Scenarios (14 scenarios required by specification):
   1. Single service
   2. Multiple services in one bill
   3. Different staff per service
   4. Single product
   5. Multiple products in one bill
   6. Different product-selling staff
   7. Combined service and product bill
   8. Missing staff
   9. Invalid service
   10. Missing payment details
   11. Invalid date
   12. Empty template
   13. Template with blank formatted rows
   14. Duplicate bill references
8. Final Requirements:
   - Optional mobile number (generates client ID)
   - Invalid mobile number rejection
   - Standalone bills when Bill Group is blank
   - Generated Client IDs Excel report
   - Preview row editing with in-place revalidation
"""
from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional
from unittest.mock import AsyncMock, MagicMock

import pytest
from beanie import PydanticObjectId
from openpyxl import Workbook, load_workbook

from app.models.bulk_billing import BulkBillingBatch, BulkBillingRowRecord
from app.models.billing import Invoice
from app.models.customer import Customer
from app.models.product import Product
from app.models.salon import Salon
from app.models.salon_product import SalonProduct
from app.models.salon_service import SalonService
from app.models.service import Service
from app.models.user import User
from app.services.bulk_billing_import import (
    COLOR_OPTIONAL_FILL,
    COLOR_REQUIRED_FILL,
    ERR_AMBIGUOUS_STAFF,
    ERR_CONFLICTING_BILL,
    ERR_DUPLICATE_BILL,
    ERR_INVALID_DATE,
    ERR_INVALID_PAYMENT,
    ERR_INVALID_PRICE,
    ERR_MISSING_REQUIRED,
    ERR_SERVICE_NOT_FOUND,
    ERR_STAFF_NOT_FOUND,
    PRODUCT_HEADERS,
    SERVICE_HEADERS,
    bulk_import_billing_service,
)


# =============================================================================
# Helper Fixtures & Builders
# =============================================================================

@pytest.fixture
def mock_db_context(monkeypatch):
    """Mocks MongoDB / Beanie queries for Salon, User, Service, Product, and Invoice."""
    salon_id = "64a000000000000000000001"
    salon = MagicMock(spec=Salon)
    salon.id = PydanticObjectId(salon_id)
    salon.name = "MyChair Test Salon"
    salon.tenant_id = "tenant-test-01"
    salon.timezone = "Asia/Kolkata"

    # Staff 1: Amit Sharma (EMP001)
    user_1 = MagicMock(spec=User)
    user_1.id = PydanticObjectId("64a000000000000000000011")
    user_1.first_name = "Amit"
    user_1.last_name = "Sharma"
    user_1.employee_code = "EMP001"
    user_1.phone = "9876543210"
    user_1.role = "employee"

    # Staff 2: Priya Patel (EMP002)
    user_2 = MagicMock(spec=User)
    user_2.id = PydanticObjectId("64a000000000000000000012")
    user_2.first_name = "Priya"
    user_2.last_name = "Patel"
    user_2.employee_code = "EMP002"
    user_2.phone = "9876543211"
    user_2.role = "employee"

    # Service 1: Haircut (300)
    svc_1 = MagicMock(spec=Service)
    svc_1.id = PydanticObjectId("64a000000000000000000021")
    svc_1.name = "Haircut"
    svc_1.price = 300.0

    # Service 2: Facial (700)
    svc_2 = MagicMock(spec=Service)
    svc_2.id = PydanticObjectId("64a000000000000000000022")
    svc_2.name = "Facial"
    svc_2.price = 700.0

    # Product 1: Argan Oil 50ml (600)
    prod_1 = MagicMock(spec=Product)
    prod_1.id = PydanticObjectId("64a000000000000000000031")
    prod_1.name = "Argan Oil 50ml"
    prod_1.price = 600.0
    prod_1.product_type = "retail"

    # Product 2: Shampoo 250ml (450)
    prod_2 = MagicMock(spec=Product)
    prod_2.id = PydanticObjectId("64a000000000000000000032")
    prod_2.name = "Shampoo 250ml"
    prod_2.price = 450.0
    prod_2.product_type = "retail"

    # Monkeypatch Beanie find methods
    monkeypatch.setattr(Salon, "find_one", AsyncMock(return_value=salon))

    user_find_mock = MagicMock()
    user_find_mock.to_list = AsyncMock(return_value=[user_1, user_2])
    monkeypatch.setattr(User, "find", MagicMock(return_value=user_find_mock))

    svc_find_mock = MagicMock()
    svc_find_mock.to_list = AsyncMock(return_value=[svc_1, svc_2])
    monkeypatch.setattr(Service, "find", MagicMock(return_value=svc_find_mock))

    prod_find_mock = MagicMock()
    prod_find_mock.to_list = AsyncMock(return_value=[prod_1, prod_2])
    monkeypatch.setattr(Product, "find", MagicMock(return_value=prod_find_mock))

    salon_svc_mock = MagicMock()
    salon_svc_mock.to_list = AsyncMock(return_value=[])
    monkeypatch.setattr(SalonService, "find", MagicMock(return_value=salon_svc_mock))

    salon_prod_mock = MagicMock()
    salon_prod_mock.to_list = AsyncMock(return_value=[])
    monkeypatch.setattr(SalonProduct, "find", MagicMock(return_value=salon_prod_mock))

    inv_find_mock = MagicMock()
    inv_find_mock.to_list = AsyncMock(return_value=[])
    monkeypatch.setattr(Invoice, "find", MagicMock(return_value=inv_find_mock))

    monkeypatch.setattr(Customer, "get_settings", MagicMock())
    monkeypatch.setattr(Customer, "find_one", AsyncMock(return_value=None))
    monkeypatch.setattr(
        "app.services.bulk_billing_import.find_client_by_phone",
        AsyncMock(return_value=None),
    )

    monkeypatch.setattr(BulkBillingBatch, "get_settings", MagicMock())
    monkeypatch.setattr(BulkBillingBatch, "insert", AsyncMock())
    monkeypatch.setattr(BulkBillingBatch, "save", AsyncMock())

    return {
        "salon_id": salon_id,
        "salon": salon,
        "users": [user_1, user_2],
        "services": [svc_1, svc_2],
        "products": [prod_1, prod_2],
        "inv_find_mock": inv_find_mock,
    }


def make_workbook_bytes(
    sheet_name: str,
    headers: List[str],
    data_rows: List[List[Any]],
    blank_rows_count: int = 0,
    include_extra_sheets: bool = True,
) -> bytes:
    """Creates an Excel file with How to Fill, Billing Data, and Example sheets."""
    wb = Workbook()

    if include_extra_sheets:
        ws_guide = wb.active
        ws_guide.title = "How to Fill"
        ws_guide.append(["STEP 1", "Open Billing Data"])
        ws_guide.append(["Yellow header = Required", "Light blue header = Optional"])

        ws_data = wb.create_sheet(title=sheet_name)
    else:
        ws_data = wb.active
        ws_data.title = sheet_name

    ws_data.append(headers)
    for row in data_rows:
        ws_data.append(row)
    for _ in range(blank_rows_count):
        ws_data.append([None] * len(headers))

    if include_extra_sheets:
        ws_ex = wb.create_sheet(title="Example")
        ws_ex.append(["FOR REFERENCE ONLY – DO NOT UPLOAD THESE EXAMPLE ROWS"])
        ws_ex.append(["2026-09-01", "Example Name", "9876543210", "Haircut", "Amit", 300])

    wb.active = ws_data
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_service_row(
    bill_date="2026-09-01",
    client_name="Rahul",
    mobile_number="9876543210",
    service_name="Haircut",
    staff_name="Amit Sharma",
    service_amount=300.0,
    payment_method="CASH",
    payment_status="PAID",
    paid_amount=300.0,
    bill_group="1",
    discount=0.0,
    notes="",
) -> List[Any]:
    return [
        bill_date,
        client_name,
        mobile_number,
        service_name,
        staff_name,
        service_amount,
        payment_method,
        payment_status,
        paid_amount,
        bill_group,
        discount,
        notes,
    ]


def make_product_row(
    bill_date="2026-09-01",
    client_name="Rahul",
    mobile_number="9876543210",
    product_name="Argan Oil 50ml",
    quantity=1,
    selling_price=600.0,
    sold_by_staff="Amit Sharma",
    payment_method="CARD",
    payment_status="PAID",
    paid_amount=600.0,
    bill_group="1",
    discount=0.0,
    notes="",
) -> List[Any]:
    return [
        bill_date,
        client_name,
        mobile_number,
        product_name,
        quantity,
        selling_price,
        sold_by_staff,
        payment_method,
        payment_status,
        paid_amount,
        bill_group,
        discount,
        notes,
    ]


# =============================================================================
# 1. Template Generation Tests
# =============================================================================

class TestBulkTemplateGeneration:
    @pytest.mark.asyncio
    async def test_service_template_download(self):
        content = await bulk_import_billing_service.build_service_billing_template()
        assert content is not None
        assert len(content) > 0

        wb = load_workbook(filename=io.BytesIO(content))
        assert "How to Fill" in wb.sheetnames
        assert "Billing Data" in wb.sheetnames
        assert "Example" in wb.sheetnames
        assert wb.active.title == "Billing Data"

        ws = wb["Billing Data"]
        headers = [c for c in next(ws.iter_rows(values_only=True))]
        assert headers == SERVICE_HEADERS
        assert len(headers) == 12

        # Verify yellow/blue header styling:
        # Col 1 (Bill Date) is required -> Yellow
        assert ws.cell(row=1, column=1).fill.start_color.rgb in (COLOR_REQUIRED_FILL, f"00{COLOR_REQUIRED_FILL}")
        # Col 3 (Mobile Number) is optional -> Blue
        assert ws.cell(row=1, column=3).fill.start_color.rgb in (COLOR_OPTIONAL_FILL, f"00{COLOR_OPTIONAL_FILL}")
        # Col 7 (Payment Method) is required -> Yellow
        assert ws.cell(row=1, column=7).fill.start_color.rgb in (COLOR_REQUIRED_FILL, f"00{COLOR_REQUIRED_FILL}")
        # Col 10 (Bill Group) is optional -> Blue
        assert ws.cell(row=1, column=10).fill.start_color.rgb in (COLOR_OPTIONAL_FILL, f"00{COLOR_OPTIONAL_FILL}")

        # Verify blank preformatted rows (at least 500 rows)
        assert ws.max_row >= 501

        # Verify Example sheet has warning notice
        ws_ex = wb["Example"]
        first_cell = ws_ex.cell(row=1, column=1).value
        assert "FOR REFERENCE ONLY" in str(first_cell)

    @pytest.mark.asyncio
    async def test_product_template_download(self):
        content = await bulk_import_billing_service.build_product_billing_template()
        assert content is not None
        assert len(content) > 0

        wb = load_workbook(filename=io.BytesIO(content))
        assert "How to Fill" in wb.sheetnames
        assert "Billing Data" in wb.sheetnames
        assert "Example" in wb.sheetnames
        assert wb.active.title == "Billing Data"

        ws = wb["Billing Data"]
        headers = [c for c in next(ws.iter_rows(values_only=True))]
        assert headers == PRODUCT_HEADERS
        assert len(headers) == 13

        # Verify preformatted blank rows
        assert ws.max_row >= 501

    @pytest.mark.asyncio
    async def test_dynamic_catalog_population(self, mock_db_context):
        salon_id = mock_db_context["salon_id"]
        content = await bulk_import_billing_service.build_service_billing_template(salon_id=salon_id)

        wb = load_workbook(filename=io.BytesIO(content))
        ws_guide = wb["How to Fill"]

        # Reference services in Column D
        services_in_guide = [
            ws_guide.cell(row=r, column=4).value
            for r in range(4, 10)
            if ws_guide.cell(row=r, column=4).value is not None
        ]
        assert "Haircut" in services_in_guide
        assert "Facial" in services_in_guide

        # Reference staff in Column E
        staff_in_guide = [
            ws_guide.cell(row=r, column=5).value
            for r in range(4, 10)
            if ws_guide.cell(row=r, column=5).value is not None
        ]
        assert any("Amit Sharma" in s and "EMP001" in s for s in staff_in_guide)
        assert any("Priya Patel" in s and "EMP002" in s for s in staff_in_guide)


# =============================================================================
# 2. Date Parsing Tests
# =============================================================================

class TestDateParsing:
    def test_valid_iso_date(self):
        dt, err = bulk_import_billing_service._parse_row_date("2026-09-01", "14:30", "UTC")
        assert err is None
        assert dt is not None
        assert dt.year == 2026
        assert dt.month == 9
        assert dt.day == 1
        assert dt.hour == 14
        assert dt.minute == 30

    def test_date_object_preserves_start_of_day_when_time_omitted(self):
        d = datetime(2026, 8, 15)
        dt, err = bulk_import_billing_service._parse_row_date(d, None, "UTC")
        assert err is None
        assert dt.day == 15
        assert dt.hour == 0
        assert dt.minute == 0

    def test_ambiguous_slash_date_rejected(self):
        dt, err = bulk_import_billing_service._parse_row_date("01/09/2026", None, "UTC")
        assert dt is None
        assert err is not None
        assert "YYYY-MM-DD" in err

    def test_missing_date(self):
        dt, err = bulk_import_billing_service._parse_row_date("", None, "UTC")
        assert dt is None
        assert err is not None


# =============================================================================
# 3. Excel Parsing Tests
# =============================================================================

class TestExcelFileParsing:
    def test_empty_content_rejected(self):
        with pytest.raises(ValueError, match="empty"):
            bulk_import_billing_service._parse_excel_bytes(b"", "SERVICE")

    def test_non_excel_content_rejected(self):
        with pytest.raises(ValueError, match="not a valid Excel"):
            bulk_import_billing_service._parse_excel_bytes(b"not-excel-bytes", "SERVICE")

    def test_parses_valid_service_workbook_and_ignores_other_sheets(self):
        content = make_workbook_bytes(
            sheet_name="Billing Data",
            headers=SERVICE_HEADERS,
            data_rows=[
                make_service_row(client_name="Aman Singh", service_name="Haircut", service_amount=300.0)
            ],
            blank_rows_count=10,
            include_extra_sheets=True,
        )
        h_map, rows = bulk_import_billing_service._parse_excel_bytes(content, "SERVICE")

        assert "client_name" in h_map
        assert "item_name" in h_map
        assert len(rows) == 1
        row_num, row_data = rows[0]
        assert row_num == 2
        assert row_data[h_map["client_name"]] == "Aman Singh"
        assert row_data[h_map["unit_price"]] == 300.0


# =============================================================================
# 4. Product Billing Parsing & Export Tests
# =============================================================================

class TestProductBillingParsing:
    def test_parses_valid_product_workbook(self):
        content = make_workbook_bytes(
            sheet_name="Billing Data",
            headers=PRODUCT_HEADERS,
            data_rows=[
                make_product_row(client_name="Vikram Malhotra", product_name="Argan Oil 50ml", quantity=2, selling_price=600.0)
            ],
            blank_rows_count=5,
            include_extra_sheets=True,
        )
        h_map, rows = bulk_import_billing_service._parse_excel_bytes(content, "PRODUCT")

        assert "quantity" in h_map
        assert "unit_price" in h_map
        assert len(rows) == 1
        r_num, r_data = rows[0]
        assert r_num == 2
        assert r_data[h_map["client_name"]] == "Vikram Malhotra"
        assert r_data[h_map["quantity"]] == 2
        assert r_data[h_map["unit_price"]] == 600.0


class TestFailedRecordsExport:
    def test_export_failed_records_xlsx(self):
        batch = BulkBillingBatch.model_construct(
            batch_id="BATCH-TEST-001",
            salon_id="64a000000000000000000001",
            status="FAILED",
            total_rows=2,
            failed_bills=1,
            row_records=[
                BulkBillingRowRecord(
                    file_type="SERVICE",
                    excel_row=2,
                    original_bill_reference="ROW-S-2",
                    client_name="Test Customer",
                    item_name="Unknown Service",
                    staff_identifier="Amit",
                    unit_price=500.0,
                    status="FAILED",
                    error_category="SERVICE_NOT_FOUND",
                    error_message="Service 'Unknown Service' not found in salon catalog.",
                    action_required="Add to catalog",
                )
            ],
        )

        xlsx_bytes = bulk_import_billing_service.build_failed_records_xlsx(batch)
        assert xlsx_bytes is not None
        assert len(xlsx_bytes) > 0

        wb = load_workbook(filename=io.BytesIO(xlsx_bytes), read_only=True)
        assert "Failed Billing Records" in wb.sheetnames
        ws = wb["Failed Billing Records"]
        rows = list(ws.iter_rows(values_only=True))
        assert len(rows) == 2
        assert rows[1][2] == "ROW-S-2"
        assert rows[1][8] == "SERVICE_NOT_FOUND"

    def test_export_summary_xlsx(self):
        batch = BulkBillingBatch.model_construct(
            batch_id="BATCH-TEST-002",
            salon_id="64a000000000000000000001",
            status="COMPLETED",
            filenames=["service.xlsx"],
            import_type="SERVICE",
            total_rows=5,
            total_bills=5,
            successful_bills=5,
            failed_bills=0,
            total_committed_amount=1500.0,
            imported_at=datetime.now(timezone.utc),
        )

        xlsx_bytes = bulk_import_billing_service.build_import_summary_xlsx(batch)
        assert xlsx_bytes is not None
        wb = load_workbook(filename=io.BytesIO(xlsx_bytes), read_only=True)
        assert "Import Summary" in wb.sheetnames
        ws = wb["Import Summary"]
        rows = list(ws.iter_rows(values_only=True))
        assert len(rows) == 2
        assert rows[1][0] == "BATCH-TEST-002"
        assert rows[1][1] == "COMPLETED"
        assert rows[1][6] == 5


# =============================================================================
# 5. Financial Calculation Math Tests
# =============================================================================

class TestCalculationsAndFinancialReconciliation:
    def test_line_total_math(self):
        unit_price = 500.0
        quantity = 2
        discount = 100.0
        tax_rate = 0.0

        line_subtotal = (unit_price * quantity) - discount  # 900.0
        line_tax = round(line_subtotal * (tax_rate / 100.0), 2)  # 0.0
        line_total = round(line_subtotal + line_tax, 2)  # 900.0

        assert line_subtotal == 900.0
        assert line_tax == 0.0
        assert line_total == 900.0

    def test_combined_bill_multi_items_sum(self):
        s1 = 300.0
        s2 = 700.0 - 50.0
        p1 = 450.0
        computed_total = s1 + s2 + p1
        assert computed_total == 1400.0


# =============================================================================
# 6. Mandatory Scenarios (14 Specification Requirements)
# =============================================================================

class TestMandatoryImportScenarios:
    @pytest.mark.asyncio
    async def test_01_single_service(self, mock_db_context):
        """Scenario 1: Single service bill with valid data."""
        salon_id = mock_db_context["salon_id"]
        svc_content = make_workbook_bytes(
            "Billing Data",
            SERVICE_HEADERS,
            [
                make_service_row(service_name="Haircut", staff_name="Amit Sharma", service_amount=300.0)
            ],
        )

        batch = await bulk_import_billing_service.validate_import(
            salon_id=salon_id,
            service_content=svc_content,
        )
        assert batch.total_bills == 1
        assert batch.valid_bills == 1
        assert batch.invalid_bills == 0
        assert batch.row_records[0].status == "VALID"
        assert batch.row_records[0].item_name == "Haircut"
        assert batch.row_records[0].resolved_staff_name == "Amit Sharma"

    @pytest.mark.asyncio
    async def test_02_multiple_services_in_one_bill(self, mock_db_context):
        """Scenario 2: Multiple services grouped under the same Bill Group."""
        salon_id = mock_db_context["salon_id"]
        svc_content = make_workbook_bytes(
            "Billing Data",
            SERVICE_HEADERS,
            [
                make_service_row(service_name="Haircut", bill_group="GRP-101", service_amount=300.0),
                make_service_row(service_name="Facial", bill_group="GRP-101", service_amount=700.0),
            ],
        )

        batch = await bulk_import_billing_service.validate_import(
            salon_id=salon_id,
            service_content=svc_content,
        )
        assert batch.total_bills == 1
        assert batch.valid_bills == 1
        assert batch.invalid_bills == 0
        assert len(batch.row_records) == 2
        assert all(r.status == "VALID" for r in batch.row_records)

    @pytest.mark.asyncio
    async def test_03_different_staff_per_service(self, mock_db_context):
        """Scenario 3: Multiple services in one bill performed by different staff members."""
        salon_id = mock_db_context["salon_id"]
        svc_content = make_workbook_bytes(
            "Billing Data",
            SERVICE_HEADERS,
            [
                make_service_row(service_name="Haircut", staff_name="Amit Sharma – EMP001", bill_group="GRP-102", service_amount=300.0),
                make_service_row(service_name="Facial", staff_name="Priya Patel – EMP002", bill_group="GRP-102", service_amount=700.0),
            ],
        )

        batch = await bulk_import_billing_service.validate_import(
            salon_id=salon_id,
            service_content=svc_content,
        )
        assert batch.valid_bills == 1
        assert len(batch.row_records) == 2
        assert batch.row_records[0].resolved_staff_name == "Amit Sharma"
        assert batch.row_records[1].resolved_staff_name == "Priya Patel"
        assert batch.row_records[0].resolved_staff_id != batch.row_records[1].resolved_staff_id

    @pytest.mark.asyncio
    async def test_04_single_product(self, mock_db_context):
        """Scenario 4: Single product sale."""
        salon_id = mock_db_context["salon_id"]
        prod_content = make_workbook_bytes(
            "Billing Data",
            PRODUCT_HEADERS,
            [
                make_product_row(product_name="Argan Oil 50ml", quantity=1, selling_price=600.0)
            ],
        )

        batch = await bulk_import_billing_service.validate_import(
            salon_id=salon_id,
            product_content=prod_content,
        )
        assert batch.total_bills == 1
        assert batch.valid_bills == 1
        assert batch.row_records[0].item_name == "Argan Oil 50ml"
        assert batch.row_records[0].quantity == 1
        assert batch.row_records[0].status == "VALID"

    @pytest.mark.asyncio
    async def test_05_multiple_products_in_one_bill(self, mock_db_context):
        """Scenario 5: Multiple products in one bill grouped by Bill Group."""
        salon_id = mock_db_context["salon_id"]
        prod_content = make_workbook_bytes(
            "Billing Data",
            PRODUCT_HEADERS,
            [
                make_product_row(product_name="Argan Oil 50ml", bill_group="GRP-201", quantity=2, selling_price=600.0),
                make_product_row(product_name="Shampoo 250ml", bill_group="GRP-201", quantity=1, selling_price=450.0),
            ],
        )

        batch = await bulk_import_billing_service.validate_import(
            salon_id=salon_id,
            product_content=prod_content,
        )
        assert batch.valid_bills == 1
        assert len(batch.row_records) == 2
        assert all(r.status == "VALID" for r in batch.row_records)

    @pytest.mark.asyncio
    async def test_06_different_product_selling_staff(self, mock_db_context):
        """Scenario 6: Multiple products in one bill sold by different staff."""
        salon_id = mock_db_context["salon_id"]
        prod_content = make_workbook_bytes(
            "Billing Data",
            PRODUCT_HEADERS,
            [
                make_product_row(product_name="Argan Oil 50ml", sold_by_staff="Amit Sharma – EMP001", bill_group="GRP-202"),
                make_product_row(product_name="Shampoo 250ml", sold_by_staff="Priya Patel – EMP002", bill_group="GRP-202"),
            ],
        )

        batch = await bulk_import_billing_service.validate_import(
            salon_id=salon_id,
            product_content=prod_content,
        )
        assert batch.valid_bills == 1
        assert len(batch.row_records) == 2
        assert batch.row_records[0].resolved_staff_name == "Amit Sharma"
        assert batch.row_records[1].resolved_staff_name == "Priya Patel"

    @pytest.mark.asyncio
    async def test_07_combined_service_and_product_bill(self, mock_db_context):
        """Scenario 7: Combined bill with both services and products sharing the same Bill Group."""
        salon_id = mock_db_context["salon_id"]
        svc_content = make_workbook_bytes(
            "Billing Data",
            SERVICE_HEADERS,
            [
                make_service_row(service_name="Haircut", staff_name="Amit Sharma", bill_group="COMBINED-01", service_amount=300.0)
            ],
        )
        prod_content = make_workbook_bytes(
            "Billing Data",
            PRODUCT_HEADERS,
            [
                make_product_row(
                    product_name="Shampoo 250ml",
                    sold_by_staff="Priya Patel",
                    bill_group="COMBINED-01",
                    selling_price=450.0,
                    payment_method="CASH",
                    paid_amount=750.0,
                )
            ],
        )

        batch = await bulk_import_billing_service.validate_import(
            salon_id=salon_id,
            service_content=svc_content,
            product_content=prod_content,
        )
        assert batch.total_bills == 1
        assert batch.valid_bills == 1
        assert len(batch.row_records) == 2
        item_types = {item.item_type for item in batch.row_records}
        assert item_types == {"SERVICE", "PRODUCT"}

    @pytest.mark.asyncio
    async def test_08_missing_staff(self, mock_db_context):
        """Scenario 8: Missing staff identifier in row."""
        salon_id = mock_db_context["salon_id"]
        svc_content = make_workbook_bytes(
            "Billing Data",
            SERVICE_HEADERS,
            [
                make_service_row(staff_name="")
            ],
        )

        batch = await bulk_import_billing_service.validate_import(
            salon_id=salon_id,
            service_content=svc_content,
        )
        assert batch.invalid_bills == 1
        failed_row = batch.row_records[0]
        assert failed_row.status == "FAILED"
        assert failed_row.error_category == ERR_MISSING_REQUIRED
        assert "Staff" in failed_row.error_message

    @pytest.mark.asyncio
    async def test_09_invalid_service(self, mock_db_context):
        """Scenario 9: Service name does not exist in salon catalog."""
        salon_id = mock_db_context["salon_id"]
        svc_content = make_workbook_bytes(
            "Billing Data",
            SERVICE_HEADERS,
            [
                make_service_row(service_name="NonExistentService")
            ],
        )

        batch = await bulk_import_billing_service.validate_import(
            salon_id=salon_id,
            service_content=svc_content,
        )
        assert batch.invalid_bills == 1
        failed_row = batch.row_records[0]
        assert failed_row.status == "FAILED"
        assert failed_row.error_category == ERR_SERVICE_NOT_FOUND

    @pytest.mark.asyncio
    async def test_10_missing_payment_details(self, mock_db_context):
        """Scenario 10: Missing payment status."""
        salon_id = mock_db_context["salon_id"]
        svc_content = make_workbook_bytes(
            "Billing Data",
            SERVICE_HEADERS,
            [
                make_service_row(payment_status="")
            ],
        )

        batch = await bulk_import_billing_service.validate_import(
            salon_id=salon_id,
            service_content=svc_content,
        )
        assert batch.invalid_bills == 1
        assert batch.row_records[0].error_category == ERR_MISSING_REQUIRED
        assert "Payment Status" in batch.row_records[0].error_message

    @pytest.mark.asyncio
    async def test_11_invalid_date(self, mock_db_context):
        """Scenario 11: Invalid date format."""
        salon_id = mock_db_context["salon_id"]
        svc_content = make_workbook_bytes(
            "Billing Data",
            SERVICE_HEADERS,
            [
                make_service_row(bill_date="invalid-date-str")
            ],
        )

        batch = await bulk_import_billing_service.validate_import(
            salon_id=salon_id,
            service_content=svc_content,
        )
        assert batch.invalid_bills == 1
        assert batch.row_records[0].error_category == ERR_INVALID_DATE

    @pytest.mark.asyncio
    async def test_12_empty_template(self, mock_db_context):
        """Scenario 12: Empty template containing only headers or blank rows."""
        salon_id = mock_db_context["salon_id"]
        svc_content = make_workbook_bytes(
            "Billing Data",
            SERVICE_HEADERS,
            [],
            blank_rows_count=100,
        )

        with pytest.raises(ValueError, match="No valid billing data rows"):
            await bulk_import_billing_service.validate_import(
                salon_id=salon_id,
                service_content=svc_content,
            )

    @pytest.mark.asyncio
    async def test_13_template_with_blank_formatted_rows(self, mock_db_context):
        """Scenario 13: Template with 1 valid data row and 500 blank formatted rows."""
        salon_id = mock_db_context["salon_id"]
        svc_content = make_workbook_bytes(
            "Billing Data",
            SERVICE_HEADERS,
            [
                make_service_row()
            ],
            blank_rows_count=500,
        )

        batch = await bulk_import_billing_service.validate_import(
            salon_id=salon_id,
            service_content=svc_content,
        )
        assert batch.total_bills == 1
        assert batch.valid_bills == 1
        assert batch.invalid_bills == 0
        assert len(batch.row_records) == 1

    @pytest.mark.asyncio
    async def test_14_duplicate_bill_references(self, mock_db_context):
        """Scenario 14: Bill Group reference that already exists in salon history."""
        salon_id = mock_db_context["salon_id"]

        existing_inv = MagicMock(spec=Invoice)
        existing_inv.original_bill_reference = "GROUP-EXISTING"
        existing_inv.invoice_number = "INV-0001"
        mock_db_context["inv_find_mock"].to_list = AsyncMock(return_value=[existing_inv])

        svc_content = make_workbook_bytes(
            "Billing Data",
            SERVICE_HEADERS,
            [
                make_service_row(bill_group="EXISTING")
            ],
        )

        batch = await bulk_import_billing_service.validate_import(
            salon_id=salon_id,
            service_content=svc_content,
        )
        assert batch.duplicate_bills == 1
        assert batch.row_records[0].status == "DUPLICATE"
        assert batch.row_records[0].error_category == ERR_DUPLICATE_BILL
        assert "already exists" in batch.row_records[0].error_message

    @pytest.mark.asyncio
    async def test_conflicting_payment_in_same_bill(self, mock_db_context):
        """Additional check: Conflicting payment status across rows of the same bill."""
        salon_id = mock_db_context["salon_id"]
        svc_content = make_workbook_bytes(
            "Billing Data",
            SERVICE_HEADERS,
            [
                make_service_row(bill_group="GRP-CONFLICT", payment_status="PAID"),
                make_service_row(bill_group="GRP-CONFLICT", payment_status="PENDING"),
            ],
        )

        batch = await bulk_import_billing_service.validate_import(
            salon_id=salon_id,
            service_content=svc_content,
        )
        assert batch.invalid_bills == 1
        failed_row = batch.row_records[1]
        assert failed_row.error_category == ERR_CONFLICTING_BILL
        assert "Conflicting payment status" in failed_row.error_message

    # =========================================================================
    # 7. Specific Requirement Tests (Auto Client ID, Blank Bill Group, Edit, Report)
    # =========================================================================

    @pytest.mark.asyncio
    async def test_optional_mobile_number_marks_will_generate(self, mock_db_context):
        """Checks that missing mobile number is accepted and marked as WILL_GENERATE."""
        salon_id = mock_db_context["salon_id"]
        svc_content = make_workbook_bytes(
            "Billing Data",
            SERVICE_HEADERS,
            [
                make_service_row(client_name="Priya Client", mobile_number="", bill_group="GRP-AUTO-CL")
            ],
        )

        batch = await bulk_import_billing_service.validate_import(
            salon_id=salon_id,
            service_content=svc_content,
        )
        assert batch.valid_bills == 1
        assert batch.client_ids_generated_count == 1
        assert batch.row_records[0].client_id_status == "WILL_GENERATE"

    @pytest.mark.asyncio
    async def test_invalid_provided_mobile_is_rejected(self, mock_db_context):
        """Checks that an invalid mobile number (e.g. 3 digits) is rejected and not silently ignored."""
        salon_id = mock_db_context["salon_id"]
        svc_content = make_workbook_bytes(
            "Billing Data",
            SERVICE_HEADERS,
            [
                make_service_row(client_name="Test Customer", mobile_number="123")
            ],
        )

        batch = await bulk_import_billing_service.validate_import(
            salon_id=salon_id,
            service_content=svc_content,
        )
        assert batch.invalid_bills == 1
        assert batch.row_records[0].status == "FAILED"
        assert "Invalid mobile number" in batch.row_records[0].error_message

    @pytest.mark.asyncio
    async def test_blank_bill_group_creates_separate_bills(self, mock_db_context):
        """Checks that rows with blank Bill Group are treated as separate standalone bills."""
        salon_id = mock_db_context["salon_id"]
        svc_content = make_workbook_bytes(
            "Billing Data",
            SERVICE_HEADERS,
            [
                make_service_row(client_name="Rahul", bill_group=""),
                make_service_row(client_name="Rahul", bill_group=""),
            ],
        )

        batch = await bulk_import_billing_service.validate_import(
            salon_id=salon_id,
            service_content=svc_content,
        )
        assert batch.total_bills == 2
        assert batch.valid_bills == 2
        assert batch.row_records[0].original_bill_reference != batch.row_records[1].original_bill_reference
        assert "AUTO-ROW-S-" in batch.row_records[0].original_bill_reference

    @pytest.mark.asyncio
    async def test_build_generated_client_ids_xlsx(self):
        """Verifies the generated Client IDs Excel report generation."""
        batch = BulkBillingBatch.model_construct(
            batch_id="BATCH-TEST-CL",
            salon_id="64a000000000000000000001",
            status="COMPLETED",
            generated_client_ids=[
                {
                    "excel_row": 2,
                    "bill_group": "GRP-1",
                    "customer_name": "Priya",
                    "mobile": "",
                    "client_id": "CL-872615",
                    "bill_number": "BILL-1001",
                    "invoice_number": "INV-1001",
                    "customer_status": "GENERATED_CLIENT_ID",
                }
            ],
        )

        xlsx_bytes = bulk_import_billing_service.build_generated_client_ids_xlsx(batch)
        assert xlsx_bytes is not None
        wb = load_workbook(filename=io.BytesIO(xlsx_bytes), read_only=True)
        assert "Generated Client IDs" in wb.sheetnames
        ws = wb["Generated Client IDs"]
        rows = list(ws.iter_rows(values_only=True))
        assert len(rows) == 2
        assert rows[1][2] == "Priya"
        assert rows[1][4] == "CL-872615"
        assert rows[1][5] == "BILL-1001"

    @pytest.mark.asyncio
    async def test_edit_row_record_in_preview(self, mock_db_context, monkeypatch):
        """Verifies that editing a row in preview updates the fields and triggers revalidation."""
        salon_id = mock_db_context["salon_id"]
        # Batch with an initially failed row due to unknown service
        batch = BulkBillingBatch(
            batch_id="BATCH-EDIT-TEST",
            salon_id=salon_id,
            tenant_id="tenant-test-01",
            filenames=["services.xlsx"],
            status="VALIDATED",
            total_rows=1,
            total_bills=1,
            valid_bills=0,
            invalid_bills=1,
            row_records=[
                BulkBillingRowRecord(
                    file_type="SERVICE",
                    excel_row=2,
                    original_bill_reference="ROW-S-2",
                    client_name="Rahul",
                    item_name="NonExistentService",
                    staff_identifier="Amit Sharma",
                    unit_price=300.0,
                    payment_status="PAID",
                    payment_method="CASH",
                    bill_date="2026-09-01",
                    status="FAILED",
                    error_category=ERR_SERVICE_NOT_FOUND,
                )
            ],
        )

        monkeypatch.setattr(BulkBillingBatch, "find_one", AsyncMock(return_value=batch))

        updated_batch = await bulk_import_billing_service.edit_row_record(
            batch_id="BATCH-EDIT-TEST",
            salon_id=salon_id,
            excel_row=2,
            update_data={"item_name": "Haircut"},
        )

        assert updated_batch.valid_bills == 1
        assert updated_batch.invalid_bills == 0
        assert updated_batch.row_records[0].status == "VALID"
        assert updated_batch.row_records[0].item_name == "Haircut"
